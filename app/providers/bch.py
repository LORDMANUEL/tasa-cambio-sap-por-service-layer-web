"""Banco Central de Honduras TCR provider.

BCH publishes a structured XLSX with daily Reference Exchange Rate (TCR)
observations. BCH is treated as an USD reference source, not as a commercial
bank buy/sell quote. For consensus purposes buy=sell=TCR is intentional.
"""
from __future__ import annotations

import hashlib
import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from zoneinfo import ZoneInfo

import requests
from openpyxl import load_workbook

from app.config import Settings
from .base import BankProvider, BankProviderError
from .http_client import secure_session

URL_TEMPLATE=(
    "https://www.bch.hn/operativos/INTL/"
    "LIBTIPO%20DE%20CAMBIO%20DE%20REFERENCIA/"
    "Resultados%20Diarios%20del%20Tipo%20de%20Cambio%20de%20Referencia%20"
    "%28TCR%29%20{year}.xlsx"
)

@dataclass(frozen=True)
class BchTcrResult:
    fetched_at: datetime
    source_url: str
    observed_date: date
    tcr: Decimal
    raw_hash: str


def _to_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text=str(value or "").strip()
    if not text:
        return None
    for fmt in ("%d/%m/%Y","%Y-%m-%d","%m/%d/%Y"):
        try:
            return datetime.strptime(text,fmt).date()
        except ValueError:
            pass
    return None


def _to_decimal(value) -> Decimal | None:
    try:
        out=Decimal(str(value).replace(",","").strip())
    except (InvalidOperation,ValueError,AttributeError):
        return None
    return out if out > 0 else None


def parse_bch_xlsx(content: bytes, settings: Settings, *, today: date | None=None, max_age_days: int=10) -> BchTcrResult:
    """Parse the latest non-future TCR from BCH's official workbook."""
    if not content:
        raise BankProviderError("BCH devolvió un archivo vacío.")
    try:
        wb=load_workbook(BytesIO(content),read_only=True,data_only=True)
    except Exception as exc:
        raise BankProviderError("BCH no devolvió un XLSX válido.") from exc
    try:
        local_today=today or datetime.now(ZoneInfo(settings.timezone)).date()
        candidates=[]
        # BCH publishes yearly workbooks with title/merged rows and changing
        # headings. Detect a real date column and a named reference-rate column
        # within the first 35 rows rather than assuming A1/B1.
        def norm(v):
            s=unicodedata.normalize("NFKD", str(v or "").upper())
            return re.sub(r"\\s+", " ", "".join(ch for ch in s if not unicodedata.combining(ch))).strip()
        def rate_col(s):
            return bool(re.search(r"\\bTCR\\b|TIPO DE CAMBIO DE REFERENCIA|TASA DE CAMBIO DE REFERENCIA",s))
        for ws in wb.worksheets:
            rows=ws.iter_rows(values_only=True)
            header=None
            for rowno,row in enumerate(rows,1):
                if rowno>35: break
                vals=[norm(v) for v in row]
                dates=[i for i,x in enumerate(vals) if re.search(r"\\bFECHA\\b|\\bDIA\\b",x)]
                rates=[i for i,x in enumerate(vals) if rate_col(x) and not re.search(r"COMPRA|VENTA|VARIACION|VAR\\.",x)]
                if dates and rates:
                    pairs=[(d,r) for d in dates for r in rates if d!=r]
                    if pairs:
                        header=(rowno,*min(pairs,key=lambda x:abs(x[0]-x[1])))
                        break
            if header is None:
                continue
            _,date_idx,rate_idx=header
            for row in rows:
                if max(date_idx,rate_idx)>=len(row): continue
                observed=_to_date(row[date_idx])
                value=_to_decimal(row[rate_idx])
                if observed and value is not None and observed<=local_today:
                    candidates.append((observed,value))
        if not candidates:
            raise BankProviderError("BCH: no se encontró una pareja Fecha/TCR válida en las hojas del XLSX.")
        observed,tcr=max(candidates,key=lambda item:item[0])
        age=(local_today-observed).days
        if age>max(1,int(max_age_days)):
            raise BankProviderError(f"BCH TCR está desactualizado: última observación {observed.isoformat()} ({age} días).")
        if not (Decimal(str(settings.usd_min)) <= tcr <= Decimal(str(settings.usd_max))):
            raise BankProviderError(f"BCH TCR fuera de rango esperado: {tcr}.")
        return BchTcrResult(
            fetched_at=datetime.now(ZoneInfo(settings.timezone)),
            source_url=URL_TEMPLATE.format(year=local_today.year),
            observed_date=observed,
            tcr=tcr,
            raw_hash=hashlib.sha256(content).hexdigest(),
        )
    finally:
        wb.close()


class BchProvider(BankProvider):
    def __init__(self,settings:Settings):
        self.settings=settings
        self.session=secure_session()

    def fetch(self)->BchTcrResult:
        local_today=datetime.now(ZoneInfo(self.settings.timezone)).date()
        url=URL_TEMPLATE.format(year=local_today.year)
        try:
            response=self.session.get(
                url,
                timeout=self.settings.request_timeout_seconds,
                verify=self.settings.tls_verify,
                allow_redirects=True,
            )
            response.raise_for_status()
            if len(response.content)>2_000_000:
                raise BankProviderError("Respuesta BCH excede el tamaño permitido.")
            result=parse_bch_xlsx(response.content,self.settings,today=local_today)
            return BchTcrResult(
                fetched_at=result.fetched_at,
                source_url=response.url or url,
                observed_date=result.observed_date,
                tcr=result.tcr,
                raw_hash=result.raw_hash,
            )
        except BankProviderError:
            raise
        except requests.RequestException as exc:
            raise BankProviderError(f"No se pudo consultar BCH: {exc.__class__.__name__}") from exc
